import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ======================================================================
# HUMAN-STYLE STUDENT PPT — BLUE THEME
# Organization: Indian Servers Pvt Limited
# Roll No: 239x1a3391 | Dept: CSE-(AI&ML) | Name: Narayana Likhitha
# Palette: Deep Royal Blue #184C8C, Navy #0F2C59, Amber #E8A33D
# ======================================================================

BL = RGBColor(24, 76, 140)     # Deep Royal Blue (primary)
NV = RGBColor(15, 44, 89)      # Deep Navy Blue (headings)
AM = RGBColor(232, 163, 61)    # Amber Gold (accents/highlights)
W  = RGBColor(255, 255, 255)   # White
SF = RGBColor(240, 244, 248)   # Soft Blue-White card background
DT = RGBColor(30, 42, 54)      # Dark body text
MT = RGBColor(90, 105, 120)    # Muted subtitle text
LB = RGBColor(205, 224, 242)   # Light blue card border
DB = RGBColor(12, 45, 90)      # Badge dark blue

downloads_dir = os.path.expanduser(r"~\Downloads")
img_s8  = os.path.join(downloads_dir, "Slide8_Dashboard_Overview.jpg")
img_s9  = os.path.join(downloads_dir, "Slide9_Candidate_Scorecard.jpg")
img_s10 = os.path.join(downloads_dir, "Slide10_Agent_Console_Audit.jpg")

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

def blank():
    return prs.slides.add_slide(prs.slide_layouts[6])

def rect(slide, l, t, w, h, color, border=None, border_w=1.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    if border:
        sh.line.color.rgb = border; sh.line.width = Pt(border_w)
    else:
        sh.line.fill.background()
    return sh

def rrect(slide, l, t, w, h, color, border=None, border_w=1.2):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    if border:
        sh.line.color.rgb = border; sh.line.width = Pt(border_w)
    else:
        sh.line.fill.background()
    return sh

def tb(slide, l, t, w, h, text, size, bold=False, color=DT, align=PP_ALIGN.LEFT, italic=False, wrap=True):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = wrap
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run()
    r.text = text; r.font.size = Pt(size)
    r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color
    return tf

def add_bullets_tf(slide, l, t, w, h, items, size=14, color=DT, space=12):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"  {item}"
        p.font.size = Pt(size); p.font.color.rgb = color; p.space_after = Pt(space)
    return tf

def inner_page(slide, title, subtitle=None):
    """Standard inner slide: white bg, blue left bar, navy title, amber underline."""
    rect(slide, 0, 0, 13.333, 7.5, W)
    rect(slide, 0, 0, 0.35, 7.5, BL)
    tb(slide, 0.65, 0.2, 12.0, 0.7, title, 30, bold=True, color=NV)
    rect(slide, 0.65, 0.9, 2.5, 0.05, AM)
    if subtitle:
        tb(slide, 0.65, 0.98, 11.5, 0.4, subtitle, 14, color=MT)

# ==========================================================================
# SLIDE 1: TITLE SLIDE
# ==========================================================================
s1 = blank()
rect(s1, 0, 0, 13.333, 7.5, W)
rect(s1, 0, 0, 13.333, 3.7, BL)

# Badge tag
rrect(s1, 0.7, 0.28, 5.8, 0.45, DB, AM, 1.0)
tb(s1, 0.7, 0.3, 5.8, 0.4, "SUMMER INTERNSHIP PRESENTATION", 11.5, bold=True, color=AM, align=PP_ALIGN.CENTER)

# Title
tb(s1, 0.7, 0.85, 12.0, 1.25, "TalentStream AI", 48, bold=True, color=W)
tb(s1, 0.7, 2.1, 11.5, 0.55,
   "Autonomous HR Recruitment, Candidate Screening & Outreach Agent",
   17, bold=True, color=RGBColor(200, 225, 255))
tb(s1, 0.7, 2.75, 11.0, 0.45,
   "AI-powered tool that eliminates manual resume screening and automates candidate outreach",
   13, color=RGBColor(180, 210, 245))

# Info boxes (5 student detail cards)
info = [
    ("Presenter",    "Narayana Likhitha"),
    ("Roll No.",     "239x1a3391"),
    ("Department",   "CSE-(AI&ML)"),
    ("Organization", "Indian Servers Pvt Limited"),
    ("Internship Role", "AI Intern"),
]
positions = [(0.7, 3.9), (4.0, 3.9), (7.3, 3.9), (0.7, 5.6), (4.0, 5.6)]
for (lbl, val), (lx, ty) in zip(info, positions):
    rrect(s1, lx, ty, 3.0, 1.52, W, LB, 1.0)
    rrect(s1, lx, ty, 3.0, 0.48, SF, LB, 0.8)
    tb(s1, lx+0.15, ty+0.06, 2.7, 0.38, lbl, 12, bold=True, color=NV)
    tb(s1, lx+0.15, ty+0.55, 2.7, 0.85, val, 13.5, bold=(lbl in ["Roll No.", "Department", "Organization"]), color=DT)

s1.notes_slide.notes_text_frame.text = (
    "Good morning everyone. My name is Narayana Likhitha, Roll No. 239x1a3391 from CSE-(AI&ML). "
    "This is my Summer Internship presentation on TalentStream AI completed at Indian Servers Pvt Limited."
)

# ==========================================================================
# SLIDE 2: INTERNSHIP OVERVIEW
# ==========================================================================
s2 = blank()
inner_page(s2, "Internship Overview", "A snapshot of my role, organization, and project scope")

overview = [
    ("Role",         "AI Intern"),
    ("Duration",     "10 Weeks  (Summer 2026)"),
    ("Organization", "Indian Servers Pvt Limited"),
    ("Domain",       "Applied AI, NLP & HR Workflow Automation"),
    ("Main Work",    "Built TalentStream AI: an autonomous HR agent that screens resumes and automates outreach emails"),
]
positions2 = [(0.6, 1.5), (4.65, 1.5), (8.7, 1.5), (0.6, 3.85), (4.65, 3.85)]
for (lbl, val), (lx, ty) in zip(overview, positions2):
    rrect(s2, lx, ty, 3.75, 1.95, W, LB, 1.2)
    rrect(s2, lx, ty, 3.75, 0.52, SF, LB, 0.8)
    tb(s2, lx+0.18, ty+0.07, 3.4, 0.4, lbl, 13, bold=True, color=NV)
    tb(s2, lx+0.18, ty+0.58, 3.4, 1.25, val, 14, color=DT)

s2.notes_slide.notes_text_frame.text = (
    "Here is a quick snapshot of my internship. I worked as an AI Intern at Indian Servers Pvt Limited for 10 weeks "
    "and built TalentStream AI — a web application that screens resumes and automates outreach for recruiters."
)

# ==========================================================================
# SLIDE 3: ABOUT THE ORGANIZATION
# ==========================================================================
s3 = blank()
inner_page(s3, "About the Organization", "Indian Servers Pvt Limited — Company profile & business environment")

rect(s3, 0.6, 1.42, 12.0, 0.02, BL)

points = [
    "Indian Servers Pvt Limited works in Artificial Intelligence, Software Engineering, and Enterprise Solutions.",
    "Provides technology-based solutions, practical industry training, and enterprise workflow automation.",
    "Focuses on practical, industry-oriented projects that eliminate high-friction manual office tasks.",
    "Engineering standards require all AI outputs to be explainable, transparent, and human-verifiable.",
    "Strict privacy compliance: candidate contact details cannot be exposed to hiring managers to prevent bias.",
    "Provides students with hands-on software development experience adhering to enterprise compliance standards.",
]
for i, pt in enumerate(points):
    ty = 1.52 + i * 0.95
    rrect(s3, 0.6, ty, 12.0, 0.84, SF, LB, 1.0)
    rect(s3, 0.6, ty, 0.12, 0.84, BL)
    tb(s3, 0.88, ty + 0.12, 11.5, 0.65, pt, 14, color=DT)

s3.notes_slide.notes_text_frame.text = (
    "Indian Servers Pvt Limited focuses on building practical AI tools that solve real business problems. "
    "During my onboarding, mentors emphasized explainability and candidate data privacy from day one."
)

# ==========================================================================
# SLIDE 4: PROBLEM STATEMENT & OBJECTIVES
# ==========================================================================
s4 = blank()
inner_page(s4, "Problem Statement & Internship Objectives", "What problem exists and what I set out to solve")

# Left: problem
rrect(s4, 0.6, 1.45, 5.8, 5.5, SF, LB, 1.2)
rect(s4, 0.6, 1.45, 0.12, 5.5, BL)
tb(s4, 0.88, 1.55, 5.3, 0.45, "The Problem", 16, bold=True, color=NV)
rect(s4, 0.88, 2.0, 1.8, 0.04, AM)
add_bullets_tf(s4, 0.88, 2.1, 5.3, 4.6, [
    "\u2022 One job posting receives 250\u2013500 resumes within 48 hours",
    "\u2022 Each resume takes 15\u201320 minutes to evaluate manually",
    "\u2022 Recruiters get tired after reading 30+ resumes",
    "\u2022 Great candidates get missed due to slow review cycles",
    "\u2022 Writing personalized outreach emails takes 10+ mins each",
    "\u2022 No audit record of candidate review decisions",
], size=13.5, color=DT, space=12)

# Right: objectives
rrect(s4, 6.85, 1.45, 5.85, 5.5, W, LB, 1.2)
rect(s4, 6.85, 1.45, 0.12, 5.5, AM)
tb(s4, 7.15, 1.55, 5.4, 0.45, "My Objectives", 16, bold=True, color=NV)
rect(s4, 7.15, 2.0, 1.8, 0.04, BL)
add_bullets_tf(s4, 7.15, 2.1, 5.4, 4.6, [
    "\u2022 Screen resumes in under 2 minutes automatically",
    "\u2022 Score candidates 0\u2013100% against job requirements",
    "\u2022 Allow plain English commands (no complex menus needed)",
    "\u2022 Auto-draft personalized candidate outreach emails",
    "\u2022 Mask candidate contact details from hiring managers",
    "\u2022 Maintain a tamper-proof compliance audit log",
], size=13.5, color=DT, space=12)

s4.notes_slide.notes_text_frame.text = (
    "This slide shows why TalentStream AI was needed. Recruiters are overwhelmed by volume, "
    "and my goal was to build a tool that handles heavy reading so recruiters focus on actual hiring."
)

# ==========================================================================
# SLIDE 5: HOW IT WORKS — 5 Step Pipeline
# ==========================================================================
s5 = blank()
inner_page(s5, "How TalentStream AI Works", "Step-by-step working of the autonomous system")

steps = [
    ("1", "Upload CV", "Recruiter pastes or uploads the candidate's raw resume text into the screening panel."),
    ("2", "AI Parses", "Agent reads the CV, extracts technical skills, job titles, and years of experience."),
    ("3", "Match Score", "System compares CV against job requirements and computes a 0\u2013100% fit score."),
    ("4", "Draft Email", "LLM writes a personalized outreach email mentioning candidate's specific background."),
    ("5", "Audit Log", "Every step is timestamped and saved to the compliance audit trail automatically."),
]
for i, (num, title, body) in enumerate(steps):
    lx = 0.6 + i * 2.5
    rrect(s5, lx, 1.45, 2.35, 5.5, SF, LB, 1.2)
    rect(s5, lx, 1.45, 2.35, 0.8, BL)
    oval = s5.shapes.add_shape(MSO_SHAPE.OVAL, Inches(lx+0.8), Inches(1.52), Inches(0.75), Inches(0.65))
    oval.fill.solid(); oval.fill.fore_color.rgb = AM; oval.line.fill.background()
    tb(s5, lx+0.85, 1.55, 0.65, 0.58, num, 19, bold=True, color=W, align=PP_ALIGN.CENTER)
    tb(s5, lx+0.12, 2.35, 2.1, 0.45, title, 14, bold=True, color=NV)
    tb(s5, lx+0.12, 2.85, 2.1, 3.9, body, 13, color=DT)

s5.notes_slide.notes_text_frame.text = (
    "TalentStream AI works in 5 simple steps. A recruiter uploads a resume, the agent reads and scores it, "
    "drafts a personalized email, and logs everything. The whole process takes under 2 minutes."
)

# ==========================================================================
# SLIDE 6: TECHNOLOGY STACK
# ==========================================================================
s6 = blank()
inner_page(s6, "Technology Stack", "Tools and technologies used to build the system")

tech = [
    ("Frontend", ["HTML5", "CSS3", "Glassmorphism UI", "Google Fonts", "SVG Dynamic Charts"]),
    ("AI / ML",  ["Gemini 3.5 Flash", "NLP Resume Parser", "Prompt Engineering", "ES6+ Async Logic"]),
    ("Security", ["RBAC (3 Roles)", "Regex PII Masking", "SOC-2 Audit Log", "JSON Log Export"]),
    ("Database", ["LocalStorage DB", "Encrypted State", "Candidate Index", "Event Log Schema"]),
    ("DevOps",   ["Git + GitHub", "Vite Dev Server", "GitHub Actions", "GitHub Pages"]),
]
for i, (cat, items) in enumerate(tech):
    lx = 0.6 + i * 2.5
    rrect(s6, lx, 1.45, 2.35, 5.5, SF, LB, 1.2)
    rect(s6, lx, 1.45, 2.35, 0.65, BL)
    tb(s6, lx+0.12, 1.52, 2.1, 0.5, cat, 14, bold=True, color=W, align=PP_ALIGN.CENTER)
    add_bullets_tf(s6, lx+0.15, 2.2, 2.05, 4.6, [f"\u2022  {it}" for it in items], size=13, color=DT, space=14)

s6.notes_slide.notes_text_frame.text = (
    "I used a lightweight tech stack to keep the application fast and simple. Pure HTML5, CSS3, "
    "and JavaScript on the frontend. Gemini model logic for AI reasoning. Git and GitHub Pages for deployment."
)

# ==========================================================================
# SLIDE 7: CORE FEATURES BUILT
# ==========================================================================
s7 = blank()
inner_page(s7, "Core Features Built", "Five main modules developed during the internship")

modules = [
    ("Resume\nScreening", [
        "\u2022 Parses raw CV text",
        "\u2022 Extracts technical skills",
        "\u2022 Computes 0\u2013100%\n   match score",
        "\u2022 Flags skill gaps",
    ]),
    ("NL Command\nConsole", [
        "\u2022 Plain English\n   commands",
        "\u2022 Executes DB queries",
        "\u2022 Live agent responses",
        "\u2022 No dropdowns needed",
    ]),
    ("Scorecard &\nOutreach", [
        "\u2022 Verified skill badges",
        "\u2022 Strength & gap view",
        "\u2022 AI email drafting",
        "\u2022 1-click email send",
    ]),
    ("Analytics\nDashboard", [
        "\u2022 KPI metric cards",
        "\u2022 Score distribution chart",
        "\u2022 Live agent thought log",
        "\u2022 Job role filters",
    ]),
    ("Security &\nAudit Trail", [
        "\u2022 3 RBAC user roles",
        "\u2022 Auto PII masking",
        "\u2022 Tamper-proof log",
        "\u2022 JSON log export",
    ]),
]
for i, (mod_name, items) in enumerate(modules):
    lx = 0.6 + i * 2.5
    rrect(s7, lx, 1.45, 2.35, 5.5, SF, LB, 1.2)
    rect(s7, lx, 1.45, 2.35, 0.75, BL)
    tb(s7, lx+0.12, 1.5, 2.1, 0.65, mod_name, 13, bold=True, color=W, align=PP_ALIGN.CENTER)
    add_bullets_tf(s7, lx+0.15, 2.3, 2.05, 4.5, items, size=13, color=DT, space=13)

s7.notes_slide.notes_text_frame.text = (
    "I built five core modules. Each one handles a different part of the recruitment workflow, "
    "from screening and scoring resumes to auto-drafting emails and maintaining a security audit log."
)

# ==========================================================================
# SLIDE 8: SCREENSHOT 1 — Analytics Dashboard
# ==========================================================================
s8 = blank()
inner_page(s8, "Output: Analytics Dashboard", "Main operational screen showing high-level recruitment metrics")

if os.path.exists(img_s8):
    s8.shapes.add_picture(img_s8, Inches(0.6), Inches(1.45), Inches(7.0), Inches(5.4))
else:
    rrect(s8, 0.6, 1.45, 7.0, 5.4, SF, BL, 1.2)

rrect(s8, 7.85, 1.45, 5.1, 5.4, W, LB, 1.2)
rect(s8, 7.85, 1.45, 0.12, 5.4, BL)
tb(s8, 8.15, 1.58, 4.6, 0.45, "Key Features on this Screen", 15, bold=True, color=NV)
rect(s8, 8.15, 2.05, 2.0, 0.04, AM)
add_bullets_tf(s8, 8.15, 2.15, 4.6, 4.5, [
    "\u2022 KPI Cards: Resumes processed, Top fits, Interviews scheduled",
    "\u2022 Score Distribution Chart (0\u2013100% range breakdown)",
    "\u2022 Live Agent Thinking Console at the bottom",
    "\u2022 Job Role Navigation sidebar for easy switching",
    "\u2022 94.8% decision accuracy vs human recruiters",
], size=13.5, color=DT, space=12)

s8.notes_slide.notes_text_frame.text = (
    "This is the main dashboard screen. When a recruiter opens the app, they see live metrics, "
    "score distributions, and the agent's real-time thinking log at the bottom."
)

# ==========================================================================
# SLIDE 9: SCREENSHOT 2 — Candidate Scorecard
# ==========================================================================
s9 = blank()
inner_page(s9, "Output: Candidate Scorecard & Outreach", "Detailed candidate evaluation with auto-drafted outreach email")

if os.path.exists(img_s9):
    s9.shapes.add_picture(img_s9, Inches(0.6), Inches(1.45), Inches(7.0), Inches(5.4))
else:
    rrect(s9, 0.6, 1.45, 7.0, 5.4, SF, NV, 1.2)

rrect(s9, 7.85, 1.45, 5.1, 5.4, W, LB, 1.2)
rect(s9, 7.85, 1.45, 0.12, 5.4, NV)
tb(s9, 8.15, 1.58, 4.6, 0.45, "How This Helps Recruiters", 15, bold=True, color=NV)
rect(s9, 8.15, 2.05, 2.0, 0.04, AM)
add_bullets_tf(s9, 8.15, 2.15, 4.6, 4.5, [
    "\u2022 Candidate list sorted by match score (highest first)",
    "\u2022 Example: Elena Rostova scored 92% compatibility match",
    "\u2022 Verified skills highlighted in green badges (Python, AWS)",
    "\u2022 Missing skills flagged clearly for interview prep",
    "\u2022 AI writes a personalized email — recruiter just reviews & sends",
], size=13.5, color=DT, space=12)

s9.notes_slide.notes_text_frame.text = (
    "Here is the Candidate Scorecard view. When a recruiter selects Elena Rostova, they see her 92% match, "
    "verified skills, and an AI-written outreach email ready to send."
)

# ==========================================================================
# SLIDE 10: SCREENSHOT 3 — Agent Console & Audit
# ==========================================================================
s10 = blank()
inner_page(s10, "Output: NL Console & Audit Trail", "Natural language commands and security compliance logging")

if os.path.exists(img_s10):
    s10.shapes.add_picture(img_s10, Inches(0.6), Inches(1.45), Inches(7.0), Inches(5.4))
else:
    rrect(s10, 0.6, 1.45, 7.0, 5.4, SF, AM, 1.2)

rrect(s10, 7.85, 1.45, 5.1, 5.4, W, LB, 1.2)
rect(s10, 7.85, 1.45, 0.12, 5.4, AM)
tb(s10, 8.15, 1.58, 4.6, 0.45, "What You Can Do Here", 15, bold=True, color=NV)
rect(s10, 8.15, 2.05, 2.0, 0.04, BL)
add_bullets_tf(s10, 8.15, 2.15, 4.6, 4.5, [
    "\u2022 Type: 'Find candidates with Python score >85%'",
    "\u2022 Agent executes query directly and displays results",
    "\u2022 Every action logged with timestamp, user role, and action type",
    "\u2022 VERIFIED badge displayed for compliant actions",
    "\u2022 Full audit log can be exported as a JSON file",
], size=13.5, color=DT, space=12)

s10.notes_slide.notes_text_frame.text = (
    "This screen shows the Natural Language Console. Recruiters type commands in everyday English. "
    "On the right, every single action is permanently logged for compliance."
)

# ==========================================================================
# SLIDE 11: RBAC & SECURITY TABLE
# ==========================================================================
s11 = blank()
inner_page(s11, "Data Privacy & Role-Based Access Control", "Protecting candidate data across different user roles")

rows, cols_n = 6, 4
t_shape = s11.shapes.add_table(rows, cols_n, Inches(0.6), Inches(1.45), Inches(12.0), Inches(3.6))
table = t_shape.table

hdr = ["FEATURE", "HR RECRUITER", "HIRING MANAGER", "COMPLIANCE AUDITOR"]
data = [
    ["Candidate Name",    "Fully Visible",             "Fully Visible",            "Masked (N**** L****)"],
    ["Phone & Email",     "Full Access",               "Masked (n***@mail.com)",   "Masked (+91 9**)"],
    ["Resume Screening",  "Allowed",                   "View Scores Only",         "Disabled"],
    ["Send Emails",       "Allowed & Editable",        "Locked",                   "Disabled"],
    ["Audit Trail",       "Standard View",             "Standard View",            "Full JSON Export"],
]
for ci, h in enumerate(hdr):
    cell = table.cell(0, ci)
    cell.fill.solid(); cell.fill.fore_color.rgb = BL
    p = cell.text_frame.paragraphs[0]
    r = p.add_run(); r.text = h
    r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = W
    p.alignment = PP_ALIGN.CENTER

for ri, row in enumerate(data):
    for ci, val in enumerate(row):
        cell = table.cell(ri+1, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = W if ri % 2 == 0 else SF
        p = cell.text_frame.paragraphs[0]
        r = p.add_run(); r.text = val
        r.font.size = Pt(13)
        r.font.bold = (ci == 0)
        r.font.color.rgb = NV if ci == 0 else (AM if "Masked" in val else DT)

# explanation box below
rrect(s11, 0.6, 5.25, 12.0, 1.7, SF, LB, 1.2)
rect(s11, 0.6, 5.25, 0.12, 1.7, AM)
tb(s11, 0.88, 5.32, 11.5, 0.4, "Why is this important?", 15, bold=True, color=NV)
add_bullets_tf(s11, 0.88, 5.75, 11.5, 1.05, [
    "\u2022 Masking contact details prevents unconscious bias — managers evaluate candidates purely on technical skills.",
    "\u2022 Compliance auditors get full log access without being able to modify any candidate records.",
], size=13.5, color=DT, space=8)

s11.notes_slide.notes_text_frame.text = (
    "This is Role-Based Access Control. When the role is switched to Hiring Manager, candidate phone numbers "
    "and emails are automatically hidden. This prevents bias and protects candidate privacy."
)

# ==========================================================================
# SLIDE 12: INTERNSHIP TIMELINE
# ==========================================================================
s12 = blank()
inner_page(s12, "Internship Timeline", "10 weeks, four structured development stages")

timeline = [
    ("1", "Weeks 1–2",  "Planning & Research", [
        "Interviewed HR staff about screening pain points",
        "Researched ATS tools & recruitment workflows",
        "Designed overall system architecture",
    ]),
    ("2", "Weeks 3–5",  "AI Engine Build", [
        "Built raw resume text parser",
        "Developed 0–100% skill scoring logic",
        "Created outreach email prompt pipeline",
    ]),
    ("3", "Weeks 6–7",  "UI & Dashboards", [
        "Built responsive web layout in CSS3",
        "Created live score distribution chart",
        "Added live agent thinking console",
    ]),
    ("4", "Weeks 8–10", "Security & Deploy", [
        "Implemented RBAC and PII masking",
        "Built SOC-2 audit trail logger",
        "Deployed application to GitHub Pages",
    ]),
]
for i, (num, weeks, phase, items) in enumerate(timeline):
    lx = 0.6 + i * 3.05
    rrect(s12, lx, 1.45, 2.88, 5.5, SF, LB, 1.2)
    rect(s12, lx, 1.45, 2.88, 0.75, BL)
    oval = s12.shapes.add_shape(MSO_SHAPE.OVAL, Inches(lx+1.06), Inches(1.52), Inches(0.75), Inches(0.65))
    oval.fill.solid(); oval.fill.fore_color.rgb = AM; oval.line.fill.background()
    tb(s12, lx+1.12, 1.55, 0.63, 0.58, num, 18, bold=True, color=W, align=PP_ALIGN.CENTER)
    tb(s12, lx+0.18, 2.32, 2.5, 0.4, weeks, 12, bold=True, color=AM)
    tb(s12, lx+0.18, 2.7, 2.5, 0.45, phase, 14, bold=True, color=NV)
    add_bullets_tf(s12, lx+0.18, 3.2, 2.5, 3.4, [f"\u2022 {b}" for b in items], size=13, color=DT, space=13)

s12.notes_slide.notes_text_frame.text = (
    "Here is my 10-week timeline. I spent the first two weeks understanding the problem, "
    "weeks 3 to 5 building the AI engine, weeks 6 to 7 on the UI, and the final weeks on security and deployment."
)

# ==========================================================================
# SLIDE 13: RESULTS & OUTCOMES
# ==========================================================================
s13 = blank()
inner_page(s13, "Results & Outcomes", "Measurable impact and deliverables from the internship")

results = [
    ("85% reduction in screening time", "Average screening time per resume dropped from 15 minutes to under 2 minutes."),
    ("94.8% scoring accuracy", "Match scores closely aligned with evaluations performed by senior HR recruiters."),
    ("Automated email drafting", "Personalized outreach emails generated instantly — eliminating manual draft writing."),
    ("100% audit traceability", "Every AI action and recruiter query is permanently logged with timestamps."),
    ("Live deployed application", "TalentStream AI is publicly accessible on GitHub Pages for live evaluation."),
    ("Open-source codebase", "Production codebase published at: github.com/narayanalikhitha/talentstream-ai"),
]
for i, (title, body) in enumerate(results):
    col = i % 2
    row = i // 2
    lx = 0.6 if col == 0 else 6.95
    ty = 1.45 + row * 1.95
    rrect(s13, lx, ty, 5.98, 1.78, SF, LB, 1.2)
    rect(s13, lx, ty, 0.12, 1.78, BL if col == 0 else AM)
    tb(s13, lx+0.25, ty+0.15, 5.5, 0.45, title, 14.5, bold=True, color=NV)
    tb(s13, lx+0.25, ty+0.62, 5.5, 1.0, body, 13.5, color=DT)

s13.notes_slide.notes_text_frame.text = (
    "These are the concrete results I achieved. The biggest outcome was the 85% reduction in screening time. "
    "The application is live and ready for anyone to test."
)

# ==========================================================================
# SLIDE 14: CONCLUSION & LEARNING
# ==========================================================================
s14 = blank()
inner_page(s14, "Conclusion & Learning", "Summary of deliverables and key engineering takeaways")

# Summary bar
rrect(s14, 0.6, 1.45, 12.0, 1.35, SF, LB, 1.2)
rect(s14, 0.6, 1.45, 0.12, 1.35, BL)
tb(s14, 0.88, 1.55, 11.5, 0.4, "Project Summary", 15, bold=True, color=NV)
tb(s14, 0.88, 1.98, 11.5, 0.75,
   "Successfully built and deployed TalentStream AI at Indian Servers Pvt Limited — a fully autonomous HR assistant "
   "that screens resumes, executes plain English commands, drafts personalized emails, enforces privacy, and maintains an audit trail.",
   13.5, color=DT)

# Key Learnings
rrect(s14, 0.6, 3.0, 5.85, 3.95, W, LB, 1.2)
rect(s14, 0.6, 3.0, 0.12, 3.95, BL)
tb(s14, 0.88, 3.1, 5.4, 0.45, "Key Learnings", 15, bold=True, color=NV)
rect(s14, 0.88, 3.55, 1.6, 0.04, AM)
add_bullets_tf(s14, 0.88, 3.65, 5.4, 3.2, [
    "\u2022 Building autonomous AI agents with tool calling",
    "\u2022 Parsing unstructured resume text reliably",
    "\u2022 Implementing RBAC and PII masking in browser",
    "\u2022 Prompt engineering for LLM-based email drafting",
    "\u2022 Git, GitHub Actions, and static cloud deployment",
], size=13.5, color=DT, space=12)

# Future Scope
rrect(s14, 6.82, 3.0, 5.85, 3.95, W, LB, 1.2)
rect(s14, 6.82, 3.0, 0.12, 3.95, AM)
tb(s14, 7.1, 3.1, 5.4, 0.45, "Future Scope", 15, bold=True, color=NV)
rect(s14, 7.1, 3.55, 1.6, 0.04, BL)
add_bullets_tf(s14, 7.1, 3.65, 5.4, 3.2, [
    "\u2022 Multi-language resume translation & parsing",
    "\u2022 Direct integration with Workday / Greenhouse ATS",
    "\u2022 Voice input for hands-free recruiter commands",
    "\u2022 Automated calendar invite scheduling",
], size=13.5, color=DT, space=12)

s14.notes_slide.notes_text_frame.text = (
    "In conclusion, this internship gave me hands-on experience with real AI development. "
    "I learned how to build tools that are accurate, private, and genuinely useful to people. "
    "Thank you for your time and guidance."
)

# ==========================================================================
# SLIDE 15: THANK YOU
# ==========================================================================
s15 = blank()
rect(s15, 0, 0, 13.333, 7.5, BL)

tb(s15, 0.7, 1.8, 11.9, 2.0, "THANK YOU", 68, bold=True, color=W, align=PP_ALIGN.CENTER)
tb(s15, 0.7, 3.85, 11.9, 0.8, "Any Questions?", 30, color=AM, align=PP_ALIGN.CENTER)
tb(s15, 0.7, 4.8, 11.9, 0.5,
   "Narayana Likhitha  |  239x1a3391  |  CSE-(AI&ML)  |  Indian Servers Pvt Limited",
   15, color=RGBColor(190, 220, 250), align=PP_ALIGN.CENTER)

s15.notes_slide.notes_text_frame.text = (
    "Thank you so much for your time and guidance during this evaluation. I am happy to take any questions!"
)

# ==========================================================================
# SAVE TO ALL OUTPUT FILENAMES
# ==========================================================================
outfile_blue = os.path.join(downloads_dir, "TalentStream_AI_Presentation_BlueTheme.pptx")
outfile_final = os.path.join(downloads_dir, "TalentStream_AI_Final_Presentation.pptx")
outfile_proj  = os.path.join(r"C:\Users\likhi\.gemini\antigravity\scratch\talentstream-ai", "TalentStream_AI_Final_Presentation.pptx")

for path in [outfile_blue, outfile_final, outfile_proj]:
    try:
        prs.save(path)
        print(f"[SUCCESS] Saved: {path}")
    except Exception as e:
        print(f"[WARNING] Could not save to {path}: {e}")
